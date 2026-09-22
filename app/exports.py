from __future__ import annotations

from io import BytesIO
from pathlib import Path

from app.i18n import DEFAULT_LOCALE, t


def _rows(estimate: dict, locale: str) -> list[list[str]]:
    materials_header = [t("item", locale), t("qty", locale), t("unit", locale), t("unit_price", locale), t("line_total", locale)]
    rows = [[t("materials", locale)], materials_header]
    for item in estimate["materials"]:
        rows.append([
            item["name"], item["quantity"], item["unit"],
            f'{item["unit_price"]} {item["currency"]}',
            f'{item["total"]} {item["currency"]}',
        ])
    rows.extend([[t("works", locale)], materials_header])
    for item in estimate["works"]:
        rows.append([
            item["name"], item["quantity"], item["unit"],
            f'{item["unit_price"]} {item["currency"]}',
            f'{item["total"]} {item["currency"]}',
        ])
    rows.extend([
        [], [t("materials_total", locale), f'{estimate["materials_total"]} {estimate["currency"]}'],
        [t("works_total", locale), f'{estimate["works_total"]} {estimate["currency"]}'],
        [t("total", locale), f'{estimate["total_price"]} {estimate["currency"]}'],
        [t("duration", locale), f'{estimate["estimated_duration_days"]} {t("days", locale)}'],
    ])
    return rows


def _heading_lines(estimate: dict, locale: str) -> list[str]:
    return [
        estimate["job_name"],
        f'{t("area", locale)}: {estimate["area_m2"]} m² | {t("region", locale)}: {estimate["region"]}',
        (
            f'{t("dimensions", locale)}: '
            f'{t("floor", locale)} {estimate["dimensions"]["floor_m2"]} m², '
            f'{t("wall", locale)} {estimate["dimensions"]["wall_m2"]} m², '
            f'{t("perimeter", locale)} {estimate["dimensions"]["perimeter_m"]} m'
        ),
    ]


def _register_pdf_font():
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont

    font_name = "EstimateUnicode"
    if font_name in pdfmetrics.getRegisteredFontNames():
        return font_name
    candidates = (
        "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
        "/usr/share/fonts/truetype/liberation2/LiberationSans-Regular.ttf",
        "/System/Library/Fonts/Supplemental/Arial Unicode.ttf",
        "/System/Library/Fonts/Supplemental/Arial.ttf",
    )
    font_path = next((path for path in candidates if Path(path).is_file()), None)
    if font_path is None:
        raise RuntimeError("Brak czcionki Unicode wymaganej do eksportu PDF.")
    pdfmetrics.registerFont(TTFont(font_name, font_path))
    return font_name


def make_xls(estimate: dict, locale: str = DEFAULT_LOCALE) -> BytesIO:
    import xlwt

    workbook = xlwt.Workbook(encoding="utf-8")
    sheet = workbook.add_sheet(t("estimate_title", locale))
    title_style = xlwt.easyxf("font: bold on, height 280")
    heading_style = xlwt.easyxf("font: bold on")
    sheet.write(0, 0, t("estimate_title", locale).upper(), title_style)
    sheet.write(0, 1, estimate["job_name"], title_style)
    sheet.write(1, 0, t("region", locale))
    sheet.write(1, 1, estimate["region"])
    sheet.write(2, 0, t("area", locale))
    sheet.write(2, 1, f'{estimate["area_m2"]} m²')
    heading_markers = {t("materials", locale), t("works", locale), t("item", locale), t("total", locale)}
    for row_index, row in enumerate(_rows(estimate, locale), start=4):
        for column_index, value in enumerate(row):
            style = heading_style if row and row[0] in heading_markers else xlwt.easyxf()
            sheet.write(row_index, column_index, value, style)
    for column_index, width in enumerate((52, 18, 12, 18, 20)):
        sheet.col(column_index).width = width * 256
    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return output


def make_docx(estimate: dict, locale: str = DEFAULT_LOCALE) -> BytesIO:
    from docx import Document

    document = Document()
    document.add_heading(t("estimate_title", locale), level=0)
    for index, line in enumerate(_heading_lines(estimate, locale)):
        if index == 0:
            document.add_heading(line, level=1)
        else:
            document.add_paragraph(line)
    document.add_heading(t("sequence", locale), level=2)
    for step in estimate["sequence"]:
        document.add_paragraph(step["name"], style="List Number")

    table_header = [t("item", locale), t("qty", locale), t("unit", locale), t("unit_price", locale), t("line_total", locale)]
    for title, items in ((t("materials", locale), estimate["materials"]), (t("works", locale), estimate["works"])):
        document.add_heading(title, level=2)
        table = document.add_table(rows=1, cols=5)
        table.style = "Table Grid"
        for cell, value in zip(table.rows[0].cells, table_header):
            cell.text = value
        for item in items:
            cells = table.add_row().cells
            values = [
                item["name"], item["quantity"], item["unit"],
                f'{item["unit_price"]} {item["currency"]}',
                f'{item["total"]} {item["currency"]}',
            ]
            for cell, value in zip(cells, values):
                cell.text = value

    document.add_heading(t("total", locale), level=2)
    document.add_paragraph(f'{t("materials_total", locale)}: {estimate["materials_total"]} {estimate["currency"]}')
    document.add_paragraph(f'{t("works_total", locale)}: {estimate["works_total"]} {estimate["currency"]}')
    document.add_paragraph(f'{t("total", locale)}: {estimate["total_price"]} {estimate["currency"]}')
    document.add_paragraph(f'{t("duration", locale)}: {estimate["estimated_duration_days"]} {t("days", locale)}')
    output = BytesIO()
    document.save(output)
    output.seek(0)
    return output


def make_pdf(estimate: dict, locale: str = DEFAULT_LOCALE) -> BytesIO:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    font_name = _register_pdf_font()
    output = BytesIO()
    document = SimpleDocTemplate(output, pagesize=A4, rightMargin=15 * mm, leftMargin=15 * mm)
    styles = getSampleStyleSheet()
    for style in styles.byName.values():
        style.fontName = font_name
    story = [Paragraph(t("estimate_title", locale), styles["Title"])]
    for line in _heading_lines(estimate, locale):
        story.append(Paragraph(line, styles["Heading2"] if line == estimate["job_name"] else styles["Normal"]))
    story.append(Spacer(1, 8))
    story.append(Paragraph(t("sequence", locale), styles["Heading2"]))
    for index, step in enumerate(estimate["sequence"], start=1):
        story.append(Paragraph(f"{index}. {step['name']}", styles["Normal"]))
    table_header = [t("item", locale), t("qty", locale), t("unit", locale), t("unit_price", locale), t("line_total", locale)]
    for title, items in ((t("materials", locale), estimate["materials"]), (t("works", locale), estimate["works"])):
        story.append(Paragraph(title, styles["Heading2"]))
        data = [table_header]
        data.extend([
            [item["name"], item["quantity"], item["unit"], f'{item["unit_price"]} {item["currency"]}', f'{item["total"]} {item["currency"]}']
            for item in items
        ])
        table = Table(data, repeatRows=1, colWidths=[75 * mm, 22 * mm, 18 * mm, 30 * mm, 30 * mm])
        table.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#f3e8dc")),
            ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#d6d3d1")),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("FONTSIZE", (0, 0), (-1, -1), 8),
            ("FONTNAME", (0, 0), (-1, -1), font_name),
        ]))
        story.extend([table, Spacer(1, 8)])
    story.append(Paragraph(
        f'{t("materials_total", locale)}: {estimate["materials_total"]} {estimate["currency"]} | '
        f'{t("works_total", locale)}: {estimate["works_total"]} {estimate["currency"]} | '
        f'{t("total", locale)}: {estimate["total_price"]} {estimate["currency"]} | '
        f'{t("duration", locale)}: {estimate["estimated_duration_days"]} {t("days", locale)}',
        styles["Heading2"],
    ))
    document.build(story)
    output.seek(0)
    return output
